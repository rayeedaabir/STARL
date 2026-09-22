# =============================================================================
# STARL v2 — FINAL REPORT BUILDER  (every experiment -> paper-ready tables/figures)
# =============================================================================
# The LAST notebook. Ingests every experiment's output zip and emits one master
# workbook, a Word document with all paper tables, and the consolidated figure set.
#
# NO GPU, NO torch — pure pandas/matplotlib. Set Accelerator = None.
#
# ATTACH (as Kaggle datasets; any that are missing are skipped with a warning):
#   * the aggregation output      (STARL_aggregate_outputs.zip)      — 12-model main results
#   * the multiseed outputs       (seed123_*, seed2025_* zips)       — variance
#   * JOINT_outputs.zip           — offline upper bound
#   * KD_outputs.zip              — compression
#   * CKA_outputs.zip             — mechanism
#   * EXTERNAL_outputs.zip        — Messidor-2 / Derm7pt
#   * EXTBASE_outputs.zip         — base-vs-final external decomposition
#   * EXEMPLAR_outputs.zip        — memory-efficiency frontier
#   (also fine: one dataset containing all the zips un-extracted — they get extracted)
#
# OUTPUT: STARL_FINAL_outputs.zip containing
#   xlsx/STARL_FINAL_tables.xlsx   — one sheet per paper table + a "claims" sheet
#   docx/STARL_FINAL_tables.docx   — the same tables, Word-ready to paste
#   figures/*.png (300 dpi) + STARL_FINAL_figures.pdf
#   csv/*.csv                      — every table as CSV


# ===== CELL 1 — imports, input discovery, output dirs =====
import os, glob, json, zipfile, shutil, subprocess, sys
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
for pkg, mod in (("openpyxl", "openpyxl"), ("python-docx", "docx")):
    try: __import__(mod)
    except ImportError: subprocess.run([sys.executable, "-m", "pip", "install", "-q", pkg])

INPUT_ROOT = "/kaggle/input"
OUT = "/kaggle/working/final"
for sub in ("csv", "xlsx", "docx", "figures"):
    os.makedirs(f"{OUT}/{sub}", exist_ok=True)
plt.rcParams.update({"figure.dpi": 110, "savefig.dpi": 300, "font.size": 10,
                     "axes.grid": True, "grid.alpha": 0.3, "axes.axisbelow": True})

SEARCH = [INPUT_ROOT]
_EXT = "/kaggle/working/_unzipped"
_nested = glob.glob(f"{INPUT_ROOT}/**/*outputs*.zip", recursive=True)
for z in _nested:
    try:
        with zipfile.ZipFile(z) as zf: zf.extractall(os.path.join(_EXT, os.path.basename(z)[:-4]))
    except Exception as e: print(f"  [warn] {z}: {e}")
if _nested:
    SEARCH.append(_EXT); print(f"  extracted {len(_nested)} nested zip(s)")

def find(pattern):
    hits = []
    for root in SEARCH: hits += glob.glob(f"{root}/**/{pattern}", recursive=True)
    return sorted(set(hits))

def read_concat(pattern, tag=None):
    frames = []
    for f in find(pattern):
        try:
            df = pd.read_csv(f)
            if tag: df["_src"] = os.path.basename(f)
            frames.append(df)
        except Exception as e: print(f"  [warn] {f}: {e}")
    return pd.concat(frames, ignore_index=True).drop_duplicates() if frames else pd.DataFrame()

TASKS = ["T1_APTOS", "T2_ODIR", "T3_LAG", "T4_HAM"]
MODEL_ORDER = ["resnet18", "resnet50", "resnet101", "resnet152", "vgg11", "vgg16", "vgg19",
               "alexnet", "googlenet", "swin_tiny", "coatnet_0", "kan"]
TABLES, FIGS = {}, []
def _safe(step, fn):
    try: fn(); print(f"  ok  {step}")
    except Exception as e:
        import traceback; print(f"  [warn] {step}: {e}"); traceback.print_exc()
def save_fig(fig, name):
    p = f"{OUT}/figures/{name}.png"; fig.tight_layout(); fig.savefig(p, bbox_inches="tight")
    plt.close(fig); FIGS.append(p); print(f"  fig {name}")


# ===== CELL 2 — Table 1: headline CL results with multi-seed mean ± sd =====
# single-seed (42) results for all 12 models come from the aggregation;
# seeds 123/2025 (7 models) come from the multiseed runner -> mean ± sd where available.
agg_cl   = read_concat("cl_metrics_by_strategy.csv")          # aggregation, seed 42
agg_ret  = read_concat("retention_after_T4.csv")              # aggregation, seed 42, rehearsal
multi_cl = read_concat("clmetrics_seed*.csv")                 # seeds 123 / 2025

def build_table1():
    rows = []
    # seed-42 rehearsal from the aggregation's retention table
    base = {}
    if not agg_ret.empty:
        for _, r in agg_ret.iterrows():
            base[r["model"]] = {"ACC": r.get("ACC"), "BWT": r.get("BWT"), "FM": r.get("FM")}
    for m in MODEL_ORDER:
        vals = {"ACC": [], "BWT": [], "FM": []}
        if m in base and pd.notna(base[m]["ACC"]):
            for k in vals: vals[k].append(float(base[m][k]))
        if not multi_cl.empty:
            sub = multi_cl[(multi_cl.model == m) & (multi_cl.strategy == "rehearsal")]
            for _, r in sub.iterrows():
                for k in vals: vals[k].append(float(r[k]))
        if not vals["ACC"]: continue
        row = {"model": m, "n_seeds": len(vals["ACC"])}
        for k in ("ACC", "BWT", "FM"):
            a = np.array(vals[k]); row[f"{k}_mean"] = a.mean()
            row[f"{k}_sd"] = a.std(ddof=1) if len(a) > 1 else np.nan
        rows.append(row)
    t = pd.DataFrame(rows).sort_values("ACC_mean", ascending=False)
    t["report"] = t.apply(lambda r: (f"{r.ACC_mean:.3f} ± {r.ACC_sd:.3f}" if pd.notna(r.ACC_sd)
                                     else f"{r.ACC_mean:.3f} (1 seed)"), axis=1)
    TABLES["T1_main_CL_results"] = t
_safe("Table 1 (main CL results)", build_table1)

def fig_headline():
    t = TABLES.get("T1_main_CL_results")
    if t is None or t.empty: return
    d = t.sort_values("ACC_mean")
    fig, ax = plt.subplots(figsize=(8, 5.2))
    err = d.ACC_sd.fillna(0)
    ax.barh(d.model, d.ACC_mean, xerr=err, capsize=3, color="#4a7ebb")
    ax.set_xlabel("Average accuracy after the full 4-task sequence (rehearsal)")
    ax.set_title("Continual-learning performance by architecture (mean ± sd over seeds)")
    ax.set_xlim(0, 0.85)
    save_fig(fig, "F1_headline_ranking")
_safe("Figure 1", fig_headline)


# ===== CELL 3 — Table 2: strategy ladder (naive / rehearsal / LwF / EWC / joint) =====
joint = read_concat("joint_summary_JOINT.csv")

def _strategy_long(path):
    """cl_metrics_by_strategy.csv is written by the aggregation as a PIVOT with a two-row
    header (row 0 = metric, row 1 = strategy). Reading it flat gives columns ACC, ACC.1 ...
    and no 'strategy' column, which previously raised KeyError and silently skipped both
    Table 2 and the bounds figure. Parse the two-row header and return long format."""
    raw = pd.read_csv(path, header=[0, 1])
    mcol = raw.columns[0]
    out = []
    for _, r in raw.iterrows():
        model = r[mcol]
        if pd.isna(model): continue
        for (metric, strat) in raw.columns:
            if metric in ("ACC", "BWT", "FM") and isinstance(strat, str) and not strat.startswith("Unnamed"):
                v = r[(metric, strat)]
                if pd.notna(v):
                    out.append({"model": model, "strategy": strat, "metric": metric, "value": float(v)})
    return pd.DataFrame(out)

def build_table2():
    rows = []
    files = find("cl_metrics_by_strategy.csv")
    long = pd.DataFrame()
    if files:
        try:
            long = _strategy_long(files[0])
        except Exception as e:
            print(f"  [warn] pivot parse failed ({e}); trying long format")
        if long.empty and not agg_cl.empty and "strategy" in agg_cl.columns:
            long = agg_cl.rename(columns={"ACC": "value"}).assign(metric="ACC")
    if not long.empty:
        acc = long[long.metric == "ACC"].pivot_table(index="model", columns="strategy", values="value")
        for m in acc.index:
            rows.append({"model": m, **{c: acc.loc[m, c] for c in acc.columns}})
    t = pd.DataFrame(rows)
    if not joint.empty and not t.empty:
        j = joint.set_index("model")["joint_avg_accuracy"].to_dict()
        t["joint (upper bound)"] = t.model.map(j)
    if not t.empty and "rehearsal" in t.columns and "joint (upper bound)" in t.columns:
        t["% of ceiling reached"] = (t["rehearsal"] / t["joint (upper bound)"] * 100).round(1)
    if t.empty:
        print("  [warn] strategy ladder empty — Table 2 and the bounds figure will be skipped")
    TABLES["T2_strategy_ladder"] = t

_safe("Table 2 (strategy ladder)", build_table2)

def fig_bounds():
    t = TABLES.get("T2_strategy_ladder")
    if t is None or t.empty or "joint (upper bound)" not in t.columns: return
    d = t.dropna(subset=["joint (upper bound)"]).copy()
    if d.empty: return
    x = np.arange(len(d)); w = 0.26
    fig, ax = plt.subplots(figsize=(2 + 1.6 * len(d), 4.6))
    for k, (col, lab) in enumerate([("naive", "naive (lower bound)"),
                                    ("rehearsal", "rehearsal (ours)"),
                                    ("joint (upper bound)", "joint (upper bound)")]):
        if col in d.columns: ax.bar(x + (k - 1) * w, d[col], w, label=lab)
    ax.set_xticks(x); ax.set_xticklabels(d.model, rotation=25, ha="right")
    ax.set_ylabel("average accuracy"); ax.set_ylim(0, 1)
    ax.set_title("Where rehearsal sits between the lower and upper bounds"); ax.legend(fontsize=8)
    save_fig(fig, "F2_bounds")
_safe("Figure 2", fig_bounds)


# ===== CELL 4 — Tables 3-5: KD, CKA, exemplar frontier =====
kd   = read_concat("kd_summary_KD.csv")
cka_drift = read_concat("cka_drift.csv")
cka_cons  = read_concat("cka_seen_unseen.csv")
exem = read_concat("exemplar_results.csv")

_safe("Table 3 (KD)", lambda: TABLES.__setitem__("T3_knowledge_distillation", kd))
def build_cka():
    if cka_drift.empty: return
    t = cka_drift.pivot_table(index=["strategy", "layer"], values=["cka_linear", "cka_rbf"],
                              aggfunc="mean").reset_index()
    TABLES["T4_CKA_drift"] = t
    if not cka_cons.empty:
        TABLES["T4b_CKA_seen_unseen"] = cka_cons.pivot_table(
            index="layer", columns="task", values="seen_unseen_cosine").reset_index()
_safe("Table 4 (CKA)", build_cka)

def build_exemplar():
    if exem.empty: return
    t = exem.pivot_table(index=["model", "m"], columns="selection", values="ACC").reset_index()
    if "herding" in t.columns and "random" in t.columns:
        t["herding − random"] = t["herding"] - t["random"]
    TABLES["T5_exemplar_frontier"] = t
_safe("Table 5 (exemplar)", build_exemplar)

def fig_cka():
    if cka_drift.empty: return
    fig, ax = plt.subplots(figsize=(7.5, 4.4))
    order = ["layer1", "layer2", "layer3", "layer4"]
    for strat, g in cka_drift.groupby("strategy"):
        mm = g.groupby("layer").cka_linear.mean().reindex(order)
        ax.plot(order, mm.values, marker="o", lw=2, label=strat)
    ax.set_ylim(0, 1); ax.set_ylabel("linear CKA (task's own expert vs final model)")
    ax.set_xlabel("network depth"); ax.legend()
    ax.set_title("Mechanism: rehearsal preserves deep activation paths that naive lets drift")
    save_fig(fig, "F3_CKA_mechanism")
def fig_exemplar():
    if exem.empty: return
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.6))
    for ax, met, ttl in zip(axes, ["ACC", "FM"], ["Average accuracy ↑", "Forgetting ↓"]):
        for (mn, sel), g in exem.groupby(["model", "selection"]):
            g = g.sort_values("m")
            ax.plot(g.m, g[met], marker="o", ls="-" if sel == "herding" else "--",
                    label=f"{mn}·{sel}")
        ax.set_xscale("log"); ax.set_xticks([1, 5, 20]); ax.set_xticklabels(["1", "5", "20"])
        ax.set_xlabel("exemplars stored per class (m)"); ax.set_title(ttl)
    axes[0].set_ylabel("after full sequence"); axes[1].legend(fontsize=6, ncol=2)
    fig.suptitle("Memory-efficiency frontier (solid = herding, dashed = random)")
    save_fig(fig, "F4_exemplar_frontier")
_safe("Figure 3", fig_cka); _safe("Figure 4", fig_exemplar)


# ===== CELL 5 — Tables 6-7: external validation + the base-vs-final decomposition =====
ext  = read_concat("external_metrics.csv")
extb = read_concat("extbase_diagnosis.csv")
_safe("Table 6 (external)", lambda: TABLES.__setitem__("T6_external_validation", ext))
_safe("Table 7 (external decomposition)", lambda: TABLES.__setitem__("T7_external_decomposition", extb))

def fig_external():
    if ext.empty: return
    for ds in ext.dataset.unique():
        d = ext[ext.dataset == ds].sort_values("balanced_accuracy_restricted"
                                               if "balanced_accuracy_restricted" in ext.columns
                                               else "accuracy_restricted")
        fig, ax = plt.subplots(figsize=(8, 4.4)); y = np.arange(len(d)); h = 0.38
        ax.barh(y - h/2, d["accuracy_restricted"], h, label="accuracy")
        if "balanced_accuracy_restricted" in d.columns:
            ax.barh(y + h/2, d["balanced_accuracy_restricted"], h, label="balanced accuracy")
        if "majority_class_baseline" in d.columns:
            ax.axvline(d["majority_class_baseline"].iloc[0], ls="--", c="r", lw=1.2,
                       label="majority-class baseline")
        ax.set_yticks(y); ax.set_yticklabels(d.model); ax.set_xlim(0, 1)
        ax.set_title(f"External validation — {ds} (never seen in training)"); ax.legend(fontsize=8)
        save_fig(fig, f"F5_external_{ds}")
_safe("Figure 5", fig_external)


# ===== CELL 6 — the CLAIMS sheet: every paper claim -> evidence -> required caveat =====
CLAIMS = [
 ("Leakage removed; results are trustworthy",
  "pHash dedup + frozen 70/15/15; selection on validation only",
  "Methods; all test numbers", "State the protocol explicitly — this is the reason for the redo."),
 ("Rehearsal prevents catastrophic forgetting",
  "ACC ~0.64-0.75 vs naive ~0.21; 14/14 paired seed runs favour rehearsal (+0.486 mean)",
  "Table 1, Table 2", "Only rehearsal is a true continual chain; naive/LwF/EWC are single-step branches."),
 ("Rehearsal recovers most of the offline ceiling",
  "rehearsal reaches ~87-91% of the joint upper bound",
  "Table 2", "Joint was run on 4 representative models only."),
 ("Depth does NOT monotonically help retention",
  "ResNet-101 worst of the ResNets in all 3 seeds (2.0-4.8 sd below 152/18)",
  "Table 1", "Confirmed by multi-seed — report as a finding, not an anomaly."),
 ("ResNet-18 is among the best AND the most stable",
  "0.7445 ± 0.0084, tightest sd of all models",
  "Table 1", "Its lead over Swin-T is only 0.8 pooled sd — say 'among the best', NOT 'the best'."),
 ("Transformers are less stable under continual learning",
  "sd 0.024-0.028 vs ResNet-18's 0.008",
  "Table 1", "Architecture affects variance, not only mean."),
 ("KAN's spline head neither helps nor hurts forgetting",
  "KAN ACC == ResNet-18 ACC (same backbone)",
  "Table 1", "Expected given the shared backbone; also a pipeline sanity check. Single seed."),
 ("Rehearsal works by preserving deep activation paths",
  "CKA(own expert vs final) rehearsal >> naive at every layer; gap largest at layer4",
  "Table 4, Figure 3", "The mechanism result — answers the activation-path hypothesis."),
 ("Seen and unseen images of a class share activation paths",
  "train-vs-test class-centroid cosine 0.96-1.00 at every layer",
  "Table 4b", "True for BOTH strategies — confirms the hypothesis but is not strategy-discriminative."),
 ("~1 image per class recovers most of replay's benefit",
  "m=1 (14 images) recovers ~67% of the full-buffer gain; m=20 (280) reaches 92-103%",
  "Table 5, Figure 4", "Answers 'one instance per class' directly."),
 ("Herding beats random, most at the smallest budget",
  "mean advantage +0.028 (m=1) -> +0.017 (m=5) -> +0.011 (m=20), 3/4 models each budget",
  "Table 5", "Run-to-run noise is ~0.025, so cite the AGGREGATE trend, not single comparisons."),
 ("The continual model compresses ~11x with no loss",
  "MobileNetV2 student matches/slightly exceeds its ResNet teacher",
  "Table 3", "Distillation/born-again effect. NEVER claim MobileNetV2 beats ResNet as an architecture."),
 ("Skin knowledge generalises externally; DR does not",
  "Derm7pt balanced acc 0.36-0.49 (chance ~0.17); Messidor-2 0.22-0.26 (chance 0.20)",
  "Table 6", "Always show the majority-class baseline beside accuracy."),
 ("The external DR failure is DOMAIN SHIFT, not forgetting",
  "base expert keeps only ~20% of its above-chance skill externally BEFORE any CL; domain hit ~8x the CL hit",
  "Table 7", "Use the corrected decomposition, not the notebook's first verdict printout."),
 ("Forgetting hits rare, clinically urgent classes hardest",
  "in-domain balanced accuracy drops 2-4x more than raw accuracy (ResNets)",
  "Table 7", "Strong medical framing; invisible if only accuracy is reported."),
]
def build_claims():
    TABLES["T0_claims_and_caveats"] = pd.DataFrame(
        CLAIMS, columns=["Claim", "Evidence", "Where", "Required caveat"])
_safe("Claims sheet", build_claims)


# ===== CELL 7 — write xlsx + docx + combined PDF + bundle =====
def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    hdr = PatternFill("solid", start_color="1F3864"); hf = Font(bold=True, color="FFFFFF")
    for name, df in TABLES.items():
        if df is None or (hasattr(df, "empty") and df.empty): continue
        ws = wb.create_sheet(name[:31])
        cols = [str(c) for c in df.columns]; ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(1, c).fill = hdr; ws.cell(1, c).font = hf
            ws.cell(1, c).alignment = Alignment(horizontal="center", wrap_text=True)
        for _, r in df.iterrows():
            ws.append([round(v, 4) if isinstance(v, (float, np.floating)) else v for v in r.tolist()])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(46, len(c) + 6))
        ws.freeze_panes = "A2"
    wb.save(f"{OUT}/xlsx/STARL_FINAL_tables.xlsx")
_safe("xlsx", write_xlsx)

for name, df in TABLES.items():
    if df is not None and hasattr(df, "empty") and not df.empty:
        _safe(f"csv/{name}", lambda n=name, d=df: d.to_csv(f"{OUT}/csv/{n}.csv", index=False))

def write_docx():
    import docx
    from docx.shared import Inches, Pt
    d = docx.Document()
    d.add_heading("STARL v2 — Final Results (paper tables & figures)", level=0)
    d.add_paragraph("Auto-assembled from every experiment output. Tables are Word-ready; "
                    "figures are 300 dpi. See sheet 'T0_claims_and_caveats' for the claim→evidence→caveat map.")
    for name, df in TABLES.items():
        if df is None or df.empty: continue
        d.add_heading(name.replace("_", " "), level=2)
        show = df.head(40).round(4)
        t = d.add_table(rows=1, cols=len(show.columns)); t.style = "Light Grid Accent 1"
        for i, c in enumerate(show.columns):
            cell = t.rows[0].cells[i]; cell.text = str(c)
            for p in cell.paragraphs:
                for run in p.runs: run.font.bold = True; run.font.size = Pt(8)
        for _, r in show.iterrows():
            cells = t.add_row().cells
            for i, c in enumerate(show.columns):
                v = r[c]; cells[i].text = f"{v:.4f}" if isinstance(v, (float, np.floating)) else str(v)
                for p in cells[i].paragraphs:
                    for run in p.runs: run.font.size = Pt(8)
    d.add_heading("Figures", level=1)
    for p in FIGS:
        d.add_heading(os.path.basename(p), level=3)
        try: d.add_picture(p, width=Inches(6.2))
        except Exception: pass
    d.save(f"{OUT}/docx/STARL_FINAL_tables.docx")
_safe("docx", write_docx)

def write_pdf():
    with PdfPages(f"{OUT}/figures/STARL_FINAL_figures.pdf") as pdf:
        for p in FIGS:
            img = plt.imread(p); fig = plt.figure(figsize=(11, 8.5))
            plt.imshow(img); plt.axis("off"); pdf.savefig(fig, dpi=150); plt.close(fig)
_safe("figures pdf", write_pdf)

manifest = {"tables": list(TABLES.keys()), "figures": [os.path.basename(p) for p in FIGS],
            "sources_found": {k: len(find(v)) for k, v in {
                "aggregation": "cl_metrics_by_strategy.csv", "multiseed": "clmetrics_seed*.csv",
                "joint": "joint_summary_JOINT.csv", "kd": "kd_summary_KD.csv",
                "cka": "cka_drift.csv", "external": "external_metrics.csv",
                "extbase": "extbase_diagnosis.csv", "exemplar": "exemplar_results.csv"}.items()}}
json.dump(manifest, open(f"{OUT}/final_manifest.json", "w"), indent=2)
out_zip = "/kaggle/working/STARL_FINAL_outputs.zip"
with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as z:
    for root, _d, files in os.walk(OUT):
        for fn in files:
            full = os.path.join(root, fn); z.write(full, os.path.relpath(full, OUT))
print("\nBUNDLED ->", out_zip)
print("sources found:", manifest["sources_found"])
print(f"tables: {len([t for t in TABLES.values() if t is not None and not t.empty])} | figures: {len(FIGS)}")
missing = [k for k, v in manifest["sources_found"].items() if v == 0]
if missing: print("\n[!] NOT FOUND (attach these zips if you want those tables):", missing)
if "T1_main_CL_results" in TABLES and not TABLES["T1_main_CL_results"].empty:
    print("\n=== TABLE 1 — headline ===")
    print(TABLES["T1_main_CL_results"][["model", "n_seeds", "report", "BWT_mean", "FM_mean"]]
          .round(4).to_string(index=False))
