"""
STARL v2 — Shared Core Library
==============================
Everything the per-task training notebooks need, in one importable module.

Key idea of the redo: Phase 1 already wrote the UNIFIED class index (label_idx)
into every split manifest, so there is NO per-task label remapping here — datasets
yield unified indices directly, and a model's growing head is just indices 0..K-1.

Import in a Kaggle cell:
    import sys; sys.path.append('/kaggle/input/starl-core')      # wherever you attach this
    import starl_core as sc

Torch/timm/kan are imported lazily so the pure-Python utilities (metrics, logging,
bundling, manifest parsing) can be unit-tested without a GPU stack.
"""

import os, csv, json, time, copy, zipfile, datetime as dt
from collections import OrderedDict, defaultdict

import numpy as np
import pandas as pd

try:
    import torch
    import torch.nn as nn
    from torch.utils.data import Dataset, DataLoader, Subset
    from torchvision import models, transforms
    _HAS_TORCH = True
except Exception:                       # allows importing pure-Python helpers without torch
    _HAS_TORCH = False

from sklearn.model_selection import train_test_split
from sklearn.metrics import (accuracy_score, precision_score, recall_score, f1_score,
                             roc_auc_score, confusion_matrix, classification_report)

IMAGENET_MEAN = [0.485, 0.456, 0.406]
IMAGENET_STD  = [0.229, 0.224, 0.225]

TRANSFORMER_NAMES = ("swin_tiny", "swin_small", "swin_base",
                     "coatnet_0", "coatnet_1", "cvt_13", "cvt_21")

# ============================================================================= #
# CONFIG / SEEDING
# ============================================================================= #
def set_seed(seed=42, deterministic=False):
    np.random.seed(seed)
    if _HAS_TORCH:
        torch.manual_seed(seed)
        torch.cuda.manual_seed_all(seed)
        # deterministic=False lets cuDNN pick the fastest conv algorithms (big speedup
        # on fixed 224x224 inputs); RNGs are still seeded for reproducible-enough runs.
        torch.backends.cudnn.deterministic = deterministic
        torch.backends.cudnn.benchmark = not deterministic

def get_device():
    return torch.device("cuda" if _HAS_TORCH and torch.cuda.is_available() else "cpu")

# ============================================================================= #
# CLASS SPACE  (loaded from Phase-1 unified_classes.json)
# ============================================================================= #
class ClassSpace:
    """Holds the unified label space and per-task / cumulative class views."""
    def __init__(self, meta):
        self.all_classes = list(meta["all_classes"])
        self.label2idx   = {c: i for i, c in enumerate(self.all_classes)}
        self.task_classes = OrderedDict((t, list(cs)) for t, cs in meta["task_classes"].items())
        self.task_order   = list(meta["task_order"])
        # cumulative unified-index set after each task (head size grows with this)
        self._cum = OrderedDict()
        seen = []
        for t in self.task_order:
            for c in self.task_classes[t]:
                if c not in seen:
                    seen.append(c)
            self._cum[t] = list(seen)

    def num_classes_through(self, task):
        """Head size when `task` is the current task = |union of classes T1..task|."""
        return len(self._cum[task])

    def task_idx(self, task):
        """Unified indices belonging to a task (for restricted metrics/ROC-AUC)."""
        return [self.label2idx[c] for c in self.task_classes[task]]

    def prior_tasks(self, task):
        i = self.task_order.index(task)
        return self.task_order[:i]

def load_class_space(splits_dir):
    with open(os.path.join(splits_dir, "unified_classes.json")) as f:
        return ClassSpace(json.load(f))

# ============================================================================= #
# DATA  (manifest-driven; no ImageFolder, no re-splitting)
# ============================================================================= #
def read_manifest(csv_path, split=None):
    """Pure-Python: return list of (filepath, label_idx) for a split ('train'/'val'/'test')."""
    rows = []
    with open(csv_path, newline="") as f:
        for r in csv.DictReader(f):
            if split is None or r["split"] == split:
                rows.append((r["filepath"], int(r["label_idx"])))
    return rows

def manifest_path(splits_dir, task):
    return os.path.join(splits_dir, f"split_manifest_{task}.csv")

def build_transforms(image_size=224):
    train_tf = transforms.Compose([
        transforms.Resize((image_size, image_size)),
        transforms.RandomHorizontalFlip(),
        transforms.RandomRotation(10),
        transforms.ToTensor(),
        transforms.Normalize(IMAGENET_MEAN, IMAGENET_STD),
    ])
    eval_tf = transforms.Compose([
        transforms.Resize((image_size, image_size)),
        transforms.ToTensor(),
        transforms.Normalize(IMAGENET_MEAN, IMAGENET_STD),
    ])
    return train_tf, eval_tf

if _HAS_TORCH:
    from PIL import Image

    import hashlib as _hashlib

    class ManifestDataset(Dataset):
        """Reads (filepath, unified_label_idx) rows; returns (image_tensor, label_idx).
        If cache_dir is set, a 224px copy of each image is written on first use and reused
        every epoch/model afterwards - large source JPEGs are decoded ONCE, not per epoch."""
        def __init__(self, rows, transform, path_remap=None, cache_dir=None, cache_size=224):
            self.rows = rows
            self.transform = transform
            self.path_remap = path_remap
            self.cache_dir = cache_dir
            self.cache_size = cache_size
            if cache_dir:
                os.makedirs(cache_dir, exist_ok=True)
        def __len__(self):
            return len(self.rows)
        def _cached_path(self, path):
            key = _hashlib.md5(path.encode()).hexdigest() + ".jpg"
            cpath = os.path.join(self.cache_dir, key)
            if not os.path.exists(cpath):
                src = self.path_remap(path) if self.path_remap else path
                im = Image.open(src).convert("RGB").resize((self.cache_size, self.cache_size))
                tmp = cpath + f".{os.getpid()}.tmp"
                im.save(tmp, "JPEG", quality=95); os.replace(tmp, cpath)
            return cpath
        def __getitem__(self, i):
            path, label = self.rows[i]
            if self.cache_dir:
                img = Image.open(self._cached_path(path)).convert("RGB")
            else:
                src = self.path_remap(path) if self.path_remap else path
                img = Image.open(src).convert("RGB")
            return self.transform(img), label

def _loader(rows, transform, batch_size, shuffle, workers, path_remap=None, cache_dir=None):
    ds = ManifestDataset(rows, transform, path_remap, cache_dir)
    return DataLoader(ds, batch_size=batch_size, shuffle=shuffle, num_workers=workers,
                      persistent_workers=(workers > 0), pin_memory=_HAS_TORCH)

def build_task_loaders(splits_dir, task, batch_size=16, image_size=224, workers=2,
                       path_remap=None, cache_dir=None):
    """Return (train_loader, val_loader, test_loader) for one task from its frozen manifest."""
    train_tf, eval_tf = build_transforms(image_size)
    mp = manifest_path(splits_dir, task)
    tr = _loader(read_manifest(mp, "train"), train_tf, batch_size, True,  workers, path_remap, cache_dir)
    va = _loader(read_manifest(mp, "val"),   eval_tf,  batch_size, False, workers, path_remap, cache_dir)
    te = _loader(read_manifest(mp, "test"),  eval_tf,  batch_size, False, workers, path_remap, cache_dir)
    return tr, va, te

def build_rehearsal_loader(splits_dir, task, fraction=0.10, reh_batch=4, image_size=224,
                           workers=2, seed=42, path_remap=None):
    """Stratified `fraction` of a prior task's TRAIN split, as a small-batch replay loader."""
    train_tf, _ = build_transforms(image_size)
    rows = read_manifest(manifest_path(splits_dir, task), "train")
    labels = [l for _, l in rows]
    n = max(len(set(labels)), int(round(len(rows) * fraction)))
    _, counts = np.unique(labels, return_counts=True)
    strat = labels if np.all(counts >= 2) else None
    idx, _ = train_test_split(range(len(rows)), train_size=min(n, len(rows) - 1),
                              random_state=seed, stratify=strat)
    sub = [rows[i] for i in idx]
    return _loader(sub, train_tf, reh_batch, True, workers, path_remap), len(sub)

def build_joint_train_loader(splits_dir, tasks, batch_size=16, image_size=224, workers=2,
                             path_remap=None):
    """Concatenate the TRAIN splits of several tasks (for the joint/offline upper bound)."""
    train_tf, _ = build_transforms(image_size)
    rows = []
    for t in tasks:
        rows += read_manifest(manifest_path(splits_dir, t), "train")
    return _loader(rows, train_tf, batch_size, True, workers, path_remap), len(rows)

# ============================================================================= #
# MODELS  (reuses the project's proven constructors, generalised)
# ============================================================================= #
_NNBase = nn.Module if _HAS_TORCH else object

class TimmTransformerWrapper(_NNBase):
    def __init__(self, backbone, num_ftrs, num_classes):
        super().__init__()
        self.backbone = backbone
        self.head = nn.Sequential(nn.Linear(num_ftrs, 512), nn.ReLU(), nn.Dropout(0.2),
                                  nn.Linear(512, num_classes))
    def forward(self, x):
        x = self.backbone.forward_features(x)
        if   x.dim() == 3: x = x.mean(dim=1)
        elif x.dim() == 4: x = x.mean(dim=[1, 2]) if x.shape[-1] > x.shape[1] else x.mean(dim=[2, 3])
        return self.head(x)

class KANWrapper(_NNBase):
    def __init__(self, backbone, kan_head):
        super().__init__()
        self.backbone = backbone
        self.kan_head = kan_head
    def forward(self, x):
        return self.kan_head(self.backbone(x))

def _head_module(model, name):
    """Return the classifier submodule for a given architecture (for freezing/probing)."""
    if name in TRANSFORMER_NAMES: return model.head
    if name == "kan":            return model.kan_head
    if "resnet" in name or name == "googlenet": return model.fc
    if "vgg" in name or "alexnet" in name:      return model.classifier[6]
    raise ValueError(name)

def _build_googlenet(weights=None):
    """torchvision forbids aux_logits=False alongside pretrained weights, so build with
    defaults then disable the aux heads -> forward() returns a plain logits tensor."""
    if weights is None:
        return models.googlenet(weights=None, aux_logits=False)
    m = models.googlenet(weights=weights)   # aux_logits defaults True with pretrained
    m.aux_logits = False
    m.aux1 = None
    m.aux2 = None
    return m

def build_model(name, num_classes, device=None, seed=42, pretrained=True,
                feature_extract=True, frozen_probe=False):
    """Construct any of the 12 models with a fresh head of size `num_classes`.
    frozen_probe=True freezes the ENTIRE backbone (only the head trains)."""
    device = device or get_device()
    if name in TRANSFORMER_NAMES:
        import timm
        TIMM = {"swin_tiny": "swin_tiny_patch4_window7_224",
                "coatnet_0": "coatnet_0_rw_224"}
        backbone = timm.create_model(TIMM.get(name, name), pretrained=pretrained, num_classes=0)
        num_ftrs = backbone.num_features
        if feature_extract:
            for p in backbone.parameters(): p.requires_grad = False
            if not frozen_probe:
                last = backbone.layers[-1] if name.startswith("swin") else backbone.stages[-1]
                for p in last.parameters(): p.requires_grad = True
                if hasattr(backbone, "norm"):
                    for p in backbone.norm.parameters(): p.requires_grad = True
        model = TimmTransformerWrapper(backbone, num_ftrs, num_classes)
    elif name == "kan":
        from kan import KAN
        backbone = models.resnet18(weights=models.ResNet18_Weights.IMAGENET1K_V1 if pretrained else None)
        backbone.fc = nn.Identity()
        if feature_extract:
            for p in backbone.parameters(): p.requires_grad = False
            if not frozen_probe:
                for n, p in backbone.named_parameters():
                    if "layer4" in n or "inception5" in n: p.requires_grad = True
        model = KANWrapper(backbone, KAN(width=[512, num_classes], grid=3, k=3, seed=seed))
    else:
        W = {"resnet18": models.ResNet18_Weights.IMAGENET1K_V1,
             "resnet50": models.ResNet50_Weights.IMAGENET1K_V2,
             "resnet101": models.ResNet101_Weights.IMAGENET1K_V2,
             "resnet152": models.ResNet152_Weights.IMAGENET1K_V2,
             "alexnet": models.AlexNet_Weights.IMAGENET1K_V1,
             "vgg11": models.VGG11_BN_Weights.IMAGENET1K_V1,
             "vgg16": models.VGG16_BN_Weights.IMAGENET1K_V1,
             "vgg19": models.VGG19_BN_Weights.IMAGENET1K_V1,
             "googlenet": models.GoogLeNet_Weights.IMAGENET1K_V1}
        C = {"resnet18": models.resnet18, "resnet50": models.resnet50,
             "resnet101": models.resnet101, "resnet152": models.resnet152,
             "alexnet": models.alexnet, "vgg11": models.vgg11_bn,
             "vgg16": models.vgg16_bn, "vgg19": models.vgg19_bn,
             "googlenet": lambda weights=None: _build_googlenet(weights)}
        if name not in C: raise ValueError(f"Unknown model: {name}")
        model = C[name](weights=W[name] if pretrained else None)
        if feature_extract:
            for p in model.parameters(): p.requires_grad = False
        if "resnet" in name or name == "googlenet":
            if feature_extract and not frozen_probe:
                for n, p in model.named_parameters():
                    if "layer4" in n or "inception5" in n: p.requires_grad = True
            nf = model.fc.in_features
            model.fc = nn.Sequential(nn.Linear(nf, 512), nn.ReLU(), nn.Dropout(0.2),
                                     nn.Linear(512, num_classes))
        else:  # vgg / alexnet
            nf = model.classifier[6].in_features
            model.classifier[6] = nn.Linear(nf, num_classes)
    if frozen_probe:                        # ensure only the (fresh) head trains
        for p in _head_module(model, name).parameters(): p.requires_grad = True
    return model.to(device)

def transfer_backbone(src_model, tgt_model, name):
    """Copy backbone weights from previous-task model into new (expanded-head) model."""
    src, tgt = src_model.state_dict(), tgt_model.state_dict()
    if name in TRANSFORMER_NAMES: exclude = ("head.",)
    elif name == "kan":           exclude = ("kan_head.",)
    elif "resnet" in name or name == "googlenet": exclude = ("fc.",)
    else:                         exclude = ()
    matched = {k: v for k, v in src.items()
               if not any(k.startswith(p) for p in exclude)
               and k in tgt and tgt[k].shape == v.shape}
    tgt.update(matched); tgt_model.load_state_dict(tgt)
    return tgt_model, len(matched)

def count_params(model):
    total = sum(p.numel() for p in model.parameters())
    train = sum(p.numel() for p in model.parameters() if p.requires_grad)
    return total, train

# ============================================================================= #
# EVALUATION
# ============================================================================= #
def quick_accuracy(model, loader, device):
    """Fast accuracy over the FULL head (used for per-epoch CL tracking)."""
    model.eval(); correct = total = 0
    with torch.no_grad():
        for images, labels in loader:
            images = images.to(device); labels = labels.to(device)
            _, preds = torch.max(model(images), 1)
            correct += (preds == labels).sum().item(); total += labels.size(0)
    return correct / total if total else 0.0

def evaluate_on_task(model, loader, task_idx, all_classes, device):
    """Full metrics for one task's test set (labels already unified indices)."""
    model.eval(); ys, ps, probs = [], [], []
    n_all = len(all_classes)
    with torch.no_grad():
        for images, labels in loader:
            images = images.to(device)
            out = model(images)
            pr = torch.softmax(out, 1)
            _, preds = torch.max(out, 1)
            ys.extend(labels.numpy()); ps.extend(preds.cpu().numpy())
            probs.extend(pr.cpu().numpy())
    y, p = np.array(ys), np.array(ps)
    prob = np.array(probs)
    class_names = [all_classes[i] for i in task_idx]
    m = {"accuracy": accuracy_score(y, p),
         "precision_weighted": precision_score(y, p, average="weighted", zero_division=0),
         "recall_weighted": recall_score(y, p, average="weighted", zero_division=0),
         "f1_weighted": f1_score(y, p, average="weighted", zero_division=0),
         "f1_macro": f1_score(y, p, average="macro", zero_division=0)}
    try:
        m["roc_auc_ovr_weighted"] = (roc_auc_score(y, prob, multi_class="ovr",
                                     average="weighted", labels=task_idx)
                                     if len(np.unique(y)) > 1 else float("nan"))
    except Exception:
        m["roc_auc_ovr_weighted"] = float("nan")
    cm = confusion_matrix(y, p, labels=range(n_all))[np.ix_(task_idx, task_idx)]
    per_cls = {}
    rep = classification_report(y, p, labels=task_idx, target_names=class_names,
                                zero_division=0, output_dict=True)
    for c in class_names:
        per_cls[c] = {"f1": rep[c]["f1-score"], "recall": rep[c]["recall"],
                      "precision": rep[c]["precision"], "support": rep[c]["support"]}
    m["per_class"] = per_cls
    m["confusion_matrix"] = cm.tolist()
    m["class_names"] = class_names
    return m

# ============================================================================= #
# CONTINUAL-LEARNING METRICS  (pure Python)
# ============================================================================= #
def cl_metrics(R):
    """R: NxN nested list. R[i][j] = accuracy on task i after training task j (None if N/A).
    Returns Average Accuracy, Backward Transfer, Forgetting Measure."""
    N = len(R)
    final = [R[i][N - 1] for i in range(N) if R[i][N - 1] is not None]
    acc = float(np.mean(final)) if final else float("nan")
    bwt, fm = [], []
    for i in range(N - 1):
        if R[i][N - 1] is not None and R[i][i] is not None:
            bwt.append(R[i][N - 1] - R[i][i])
        seen = [R[i][j] for j in range(i, N) if R[i][j] is not None]
        if seen and R[i][N - 1] is not None:
            fm.append(max(seen) - R[i][N - 1])
    return {"average_accuracy": acc,
            "backward_transfer": float(np.mean(bwt)) if bwt else float("nan"),
            "forgetting_measure": float(np.mean(fm)) if fm else float("nan")}

# ============================================================================= #
# LOGGING  (crash-safe: one CSV row appended per epoch)
# ============================================================================= #
class MetricLogger:
    def __init__(self, csv_path):
        self.csv_path = csv_path
        os.makedirs(os.path.dirname(csv_path) or ".", exist_ok=True)
    def log_epoch(self, row: dict):
        write_header = not os.path.exists(self.csv_path)
        pd.DataFrame([row]).to_csv(self.csv_path, mode="a", header=write_header, index=False)

# ============================================================================= #
# TRAINING STRATEGIES
# ============================================================================= #
def _epoch_eval_row(model, model_name, experiment, epoch, ep_time, tr_loss, tr_acc,
                    eval_loaders, device):
    accs = {name: quick_accuracy(model, ld, device) for name, ld in eval_loaders.items()}
    row = {"model": model_name, "experiment": experiment, "epoch": epoch,
           "time_s": round(ep_time, 2), "train_loss": tr_loss, "train_acc": tr_acc}
    row.update({f"{name}_acc": accs[name] for name in eval_loaders})
    return row, accs

def train_naive(model, train_loader, eval_loaders, optimizer, device, epochs, model_name,
                logger=None, experiment="naive", criterion=None, select_best=False,
                select_loaders=None, best_ckpt_path=None, class_space=None,
                all_classes=None, seed=42, use_amp=True):
    """Fine-tune on one task's data only. Tracks every eval task per epoch.
    eval_loaders = what we log/report each epoch (TEST curves);
    select_loaders = what best-checkpoint selection uses (VAL) — never select on test.
    (frozen-probe = call this with a frozen_probe model.)"""
    criterion = criterion or nn.CrossEntropyLoss()
    use_amp = use_amp and device.type == "cuda"
    scaler = torch.cuda.amp.GradScaler(enabled=use_amp)
    best_avg, best_wts, history = -1.0, None, []
    for epoch in range(1, epochs + 1):
        t0 = time.time(); model.train(); run, correct, total = 0.0, 0, 0
        for images, labels in train_loader:
            images, labels = images.to(device), labels.to(device)
            optimizer.zero_grad()
            with torch.cuda.amp.autocast(enabled=use_amp):
                out = model(images); loss = criterion(out, labels)
            scaler.scale(loss).backward(); scaler.step(optimizer); scaler.update()
            run += loss.item() * images.size(0)
            _, preds = torch.max(out, 1); correct += (preds == labels).sum().item()
            total += images.size(0)
        tr_loss, tr_acc = run / total, correct / total
        row, accs = _epoch_eval_row(model, model_name, experiment, epoch, time.time() - t0,
                                    tr_loss, tr_acc, eval_loaders, device)
        history.append(row)
        if logger: logger.log_epoch(row)
        print(f"  [{model_name}/{experiment}] ep {epoch}/{epochs} "
              f"{row['time_s']:.0f}s loss {tr_loss:.4f} acc {tr_acc:.4f} | "
              + " ".join(f"{n}:{a:.3f}" for n, a in accs.items()), flush=True)
        if select_best:
            sel = ({n: quick_accuracy(model, ld, device) for n, ld in select_loaders.items()}
                   if select_loaders else accs)
            avg = float(np.mean(list(sel.values())))
            if avg > best_avg:
                best_avg, best_wts = avg, copy.deepcopy(model.state_dict())
                if best_ckpt_path:
                    save_checkpoint(model, best_ckpt_path, all_classes, seed=seed,
                                    extra={"epoch": epoch, "avg_acc": avg})
    if select_best and best_wts: model.load_state_dict(best_wts)
    return model, history

def train_rehearsal(model, new_task_loader, rehearsal_loaders, eval_loaders, optimizer,
                    device, epochs, model_name, logger=None, criterion=None,
                    select_loaders=None, best_ckpt_path=None, all_classes=None, seed=42, use_amp=True):
    """Train on the new task + replay buffers of all prior tasks (list of loaders).
    eval_loaders = per-epoch TEST tracking; select_loaders = VAL used for best checkpoint.
    Best checkpoint = highest average accuracy across all tracked (val) tasks."""
    criterion = criterion or nn.CrossEntropyLoss()
    use_amp = use_amp and device.type == "cuda"
    scaler = torch.cuda.amp.GradScaler(enabled=use_amp)
    best_avg, best_wts, history = -1.0, None, []
    for epoch in range(1, epochs + 1):
        t0 = time.time(); model.train(); run, correct, total = 0.0, 0, 0
        iters = [iter(rl) for rl in rehearsal_loaders]
        for s_images, s_labels in new_task_loader:
            imgs, lbls = [s_images], [s_labels]
            for i, rl in enumerate(rehearsal_loaders):
                try:
                    r_i, r_l = next(iters[i])
                except StopIteration:
                    iters[i] = iter(rl); r_i, r_l = next(iters[i])
                imgs.append(r_i); lbls.append(r_l)
            images = torch.cat(imgs).to(device); labels = torch.cat(lbls).to(device)
            optimizer.zero_grad()
            with torch.cuda.amp.autocast(enabled=use_amp):
                out = model(images); loss = criterion(out, labels)
            scaler.scale(loss).backward(); scaler.step(optimizer); scaler.update()
            run += loss.item() * images.size(0)
            _, preds = torch.max(out, 1); correct += (preds == labels).sum().item()
            total += images.size(0)
        tr_loss, tr_acc = run / total, correct / total
        row, accs = _epoch_eval_row(model, model_name, "rehearsal", epoch, time.time() - t0,
                                    tr_loss, tr_acc, eval_loaders, device)
        history.append(row)
        if logger: logger.log_epoch(row)
        print(f"  [{model_name}/rehearsal] ep {epoch}/{epochs} "
              f"{row['time_s']:.0f}s loss {tr_loss:.4f} acc {tr_acc:.4f} | "
              + " ".join(f"{n}:{a:.3f}" for n, a in accs.items()), flush=True)
        sel = ({n: quick_accuracy(model, ld, device) for n, ld in select_loaders.items()}
               if select_loaders else accs)
        avg = float(np.mean(list(sel.values())))
        if avg > best_avg:
            best_avg, best_wts = avg, copy.deepcopy(model.state_dict())
            if best_ckpt_path:
                save_checkpoint(model, best_ckpt_path, all_classes, seed=seed,
                                extra={"epoch": epoch, "avg_acc": avg})
    if best_wts: model.load_state_dict(best_wts)
    return model, history

# ============================================================================= #
# CHECKPOINTS + OUTPUT BUNDLE
# ============================================================================= #
def save_checkpoint(model, path, all_classes, num_classes=None, seed=42, extra=None):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    payload = {"model_state_dict": model.state_dict(),
               "all_classes": all_classes,
               "num_classes": num_classes if num_classes is not None
                               else (len(all_classes) if all_classes else None),
               "seed": seed}
    if extra: payload.update(extra)
    torch.save(payload, path)

def write_final_metrics_csv(rows, path):
    pd.DataFrame(rows).to_csv(path, index=False)

def write_perclass_csv(records, path):
    pd.DataFrame(records).to_csv(path, index=False)

def build_results_xlsx(sheets: dict, path):
    """sheets: {sheet_name: list-of-dict-rows} -> one formatted .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    hdr = PatternFill("solid", start_color="1F3864")
    hf = Font(bold=True, color="FFFFFF")
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        if not rows:
            continue
        cols = list(rows[0].keys())
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(1, c).fill = hdr; ws.cell(1, c).font = hf
            ws.cell(1, c).alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(c)) + 4))
        ws.freeze_panes = "A2"
    wb.save(path)

def build_report_docx(task, summary_rows, path, figures_dir=None, title=None):
    """Minimal Word report: summary table + any PNG figures embedded.
    Degrades gracefully if python-docx is not installed (skips the .docx)."""
    try:
        import docx
        from docx.shared import Inches
    except ImportError:
        try:
            import subprocess, sys as _sys
            subprocess.run([_sys.executable, "-m", "pip", "install", "-q", "python-docx"], check=True)
            import docx
            from docx.shared import Inches
        except Exception:
            print("  [warn] python-docx unavailable - skipping .docx (xlsx/csv still written)")
            return None
    d = docx.Document()
    d.add_heading(title or f"STARL v2 — {task} Results", level=0)
    d.add_paragraph(f"Generated {dt.datetime.utcnow().isoformat()}Z")
    if summary_rows:
        cols = list(summary_rows[0].keys())
        t = d.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols): t.rows[0].cells[i].text = str(c)
        for r in summary_rows:
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                v = r.get(c); cells[i].text = f"{v:.4f}" if isinstance(v, float) else str(v)
    if figures_dir and os.path.isdir(figures_dir):
        for fn in sorted(os.listdir(figures_dir)):
            if fn.lower().endswith(".png"):
                d.add_heading(fn, level=2)
                try: d.add_picture(os.path.join(figures_dir, fn), width=Inches(6))
                except Exception: pass
    d.save(path)

def bundle_outputs(task, working_dir, out_zip=None, run_manifest=None):
    """Zip the standard output layout: models/ csv/ xlsx/ docx/ figures/ resume/ + manifest."""
    out_zip = out_zip or os.path.join(os.path.dirname(working_dir) or ".", f"{task}_outputs.zip")
    if run_manifest:
        with open(os.path.join(working_dir, "run_manifest.json"), "w") as f:
            json.dump(run_manifest, f, indent=2)
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _dirs, files in os.walk(working_dir):
            for fn in files:
                full = os.path.join(root, fn)
                z.write(full, arcname=os.path.relpath(full, working_dir))
    return out_zip

def next_task_ckpt_name(model_name, next_task):
    """Filename the NEXT task's notebook will look for (pre-rename so no manual step)."""
    return f"expert_{next_task}_{model_name}.pth"
